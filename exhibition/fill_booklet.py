"""Fills the AIC Hub A4 booklet page template with SPIREXA project content.

Same approach as fill_poster.py: edit by shape NAME, rewrite run.text only
(never paragraph.text) to preserve template formatting. The booklet's bullet
placeholders ship with zero run formatting (like the poster's did), which
inherits an oversized default and overflows the template's tight ~0.16in
row spacing -- set_bullet() below fixes that with an explicit small size,
learned from the poster fill.
"""

from pptx import Presentation
from pptx.util import Inches, Pt

SRC = "A4_booklet_template.pptx"
OUT = "A4_Booklet_SPIREXA.pptx"
ASSETS = "/Users/utkarsh/Desktop/ml/.docxbuild/assets26"


def set_text(shape, new_text):
    """Rewrites a shape's visible text via its first run, preserving formatting."""
    if shape is None or not shape.has_text_frame:
        print(f"  WARNING: shape not found for text: {new_text[:50]!r}")
        return False
    tf = shape.text_frame
    if not tf.paragraphs or not tf.paragraphs[0].runs:
        tf.paragraphs[0].text = new_text
        return True
    tf.paragraphs[0].runs[0].text = new_text
    for r in tf.paragraphs[0].runs[1:]:
        r.text = ""
    for p in tf.paragraphs[1:]:
        for r in p.runs:
            r.text = ""
    return True


def by_name(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None


def by_text(slide, text):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip() == text:
            return sh
    return None


def set_bullet(slide, shp_name, text):
    """Like set_text, but for the empty bullet-content placeholders: these
    ship with no run at all, so they silently inherit a master default size
    too large for the template's ~0.16in row spacing. Match the 8.2pt size
    the template itself uses for adjacent pre-filled body content (e.g. the
    Current Status fields), explicitly, so it can't inherit something bigger."""
    shape = by_name(slide, shp_name)
    if not set_text(shape, text):
        return
    run = shape.text_frame.paragraphs[0].runs[0]
    run.font.size = Pt(8.2)
    run.font.name = "Calibri"


def fit_picture(slide, path, box, pad=0.04):
    """Adds an image centered inside (left, top, width, height), in inches,
    preserving its aspect ratio rather than stretching it to fill the box."""
    from PIL import Image

    left, top, w, h = box
    iw, ih = Image.open(path).size
    box_ratio, img_ratio = w / h, iw / ih
    if img_ratio > box_ratio:
        draw_w = w - 2 * pad
        draw_h = draw_w / img_ratio
    else:
        draw_h = h - 2 * pad
        draw_w = draw_h * img_ratio
    x = left + (w - draw_w) / 2
    y = top + (h - draw_h) / 2
    slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(draw_w), Inches(draw_h))


def set_para_runs(paragraph, texts):
    """Sets each existing run in a paragraph to the corresponding string in
    `texts`, leaving any <a:br/> or other structure between them untouched
    -- unlike set_text, which collapses a whole shape to one run and would
    destroy the real line break this shape uses between its two lines."""
    for run, text in zip(paragraph.runs, texts):
        run.text = text
    for run in paragraph.runs[len(texts):]:
        run.text = ""


def check(shape, option_text):
    """Marks one specific '□ Option' as '☑ Option'. Works per-paragraph (not
    via set_text on the whole shape) because unlike the poster's checkbox
    row -- one paragraph holding all options as a single spaced-out string
    -- this template's "Looking For" boxes hold each option as its own
    paragraph; rewriting the whole shape's text would collapse all of them
    into one run and blank out the rest."""
    if shape is None or not shape.has_text_frame:
        print(f"  WARNING: checkbox shape not found for: {option_text!r}")
        return
    target = f"□ {option_text}"
    checked = f"☑ {option_text}"
    for p in shape.text_frame.paragraphs:
        ptext = "".join(r.text for r in p.runs)
        if target in ptext:
            new_ptext = ptext.replace(target, checked)
            if p.runs:
                p.runs[0].text = new_ptext
                for r in p.runs[1:]:
                    r.text = ""
            else:
                p.text = new_ptext
            return
    print(f"  WARNING: checkbox option text not found: {option_text!r}")


pres = Presentation(SRC)

# Delete slide 0 (the "how to fill this template" guidelines page).
xml_slides = pres.slides._sldIdLst
slides = list(xml_slides)
xml_slides.remove(slides[0])

slide = pres.slides[0]

# ---------------------------------------------------------------- header --
# TextBox 6 is 2 paragraphs: para[0] holds 2 runs joined by a real <a:br/>
# (its two-tone "CAREER & INDUSTRY CONNECT 2026" / "STARTUP & INNOVATION
# EXHIBITION" lines), para[1] is a small italic sub-line. A plain set_text()
# with an embedded \x0b looked right when read back by python-pptx but
# actually serializes as the literal escape "_x000B_" on render -- writing
# a real line break needs the existing <a:br/> left alone and each run
# retextted individually, which is what set_para_runs does.
tb6 = by_name(slide, "TextBox 6")
set_para_runs(tb6.text_frame.paragraphs[0],
              ["AIC HUB 2026", "STUDENT PROJECT & INNOVATION EXHIBITION"])
set_para_runs(tb6.text_frame.paragraphs[1], [""])
set_text(by_name(slide, "TextBox 7"), "14th September, 2026    •    IIIT Allahabad")
set_text(by_name(slide, "TextBox 13"), "SPIREXA — AI Landslide Warning")
set_text(by_name(slide, "TextBox 14"), "AI-based early warning")
set_text(by_name(slide, "TextBox 17"), "SPIREXA")

# ------------------------------------------------------- left photo block -
fit_picture(slide, f"{ASSETS}/dashboard_full.png", (0.22, 2.0, 3.55, 2.05))
set_text(by_name(slide, "TextBox 20"), "")

# ------------------------------------------------------ right info block --
set_text(by_name(slide, "TextBox 21"), "Team SPIREXA")
set_text(by_name(slide, "TextBox 24"), "Founder(s) / Team Members: Utkarsh (Team Lead) & Team")
by_name(slide, "TextBox 24").width = Inches(3.66)
set_text(by_name(slide, "TextBox 28"),
         "□ Incubated Startup □ Student Startup □ Student Innovation □ Alumni Startup")
by_name(slide, "TextBox 28").width = Inches(3.9)
check(by_name(slide, "TextBox 28"), "Student Innovation")
# Headquarters left blank -- honest: this is a hackathon project, not a
# registered startup with a physical HQ.
set_text(by_name(slide, "TextBox 33"), "Website: github.com/UtkarshCodecom/spirexa-ner")
by_name(slide, "TextBox 33").width = Inches(3.66)
set_text(by_name(slide, "TextBox 36"), "Email: yushkarsh@gmail.com")
by_name(slide, "TextBox 36").width = Inches(3.66)
# Phone left blank -- not sharing a personal number in a public handout.

# Connector 25/34/37 are the pen-fill underlines that used to sit to the
# right of the short "Founder(s):"/"Website:"/"Email:" labels, in the blank
# space meant for handwriting. Now that those labels carry the full answer
# inline, the underlines run directly through the new text as a strikethrough
# -- Headquarters/Phone stay blank so their own underlines (31, 40) are left.
for name in ["Connector 25", "Connector 34", "Connector 37"]:
    shape = by_name(slide, name)
    shape._element.getparent().remove(shape._element)

# --------------------------------------------------- 01 PROBLEM / NEED ----
bullets_01 = [
    "NER landslides cut off roads & homes every monsoon.",
    "Warning today is reactive — only after a slope fails.",
    "NHPC Teesta-V (Aug 2024) alone: Rs 327 crore lost.",
]
for shp_name, text in zip(["TextBox 48", "TextBox 50", "TextBox 52"], bullets_01):
    set_bullet(slide, shp_name, text)

# --------------------------------------------------- 02 OUR SOLUTION ------
bullets_02 = [
    "Give any NER coordinate + date — risk % in ~10 seconds.",
    "Live GIS dashboard: 311 hotspots, recalculated daily.",
    "Alerts auto-generated in 9 regional languages.",
]
for shp_name, text in zip(["TextBox 78", "TextBox 80", "TextBox 82"], bullets_02):
    set_bullet(slide, shp_name, text)

# ---------------------------------------------- 03 KEY FEATURES/INNOVATION-
bullets_03 = [
    "Trained on 318 real NER landslides (2007-2025).",
    "Caught & fixed a road-proximity reporting bias.",
    "Physics (USGS slope-stability) + machine learning.",
]
for shp_name, text in zip(["TextBox 108", "TextBox 110", "TextBox 112"], bullets_03):
    set_bullet(slide, shp_name, text)

# ------------------------------------------------- 04 APPLICATIONS/IMPACT -
bullets_04 = [
    "District control rooms get a daily ranked watch-list.",
    "Remote villages get warnings in their own language.",
    "Field officers get offline hazard reporting.",
]
for shp_name, text in zip(["TextBox 63", "TextBox 65", "TextBox 67"], bullets_04):
    set_bullet(slide, shp_name, text)

# ------------------------------------------------------ 05 CURRENT STATUS -
set_bullet(slide, "TextBox 92", "Development Stage / TRL: TRL-4, lab-validated")
set_bullet(slide, "TextBox 94", "Prototype / Product: ML model + dashboard + Arduino sensor")
set_bullet(slide, "TextBox 96", "IP / Patent (if any): None yet")
set_bullet(slide, "TextBox 98", "Achievements (if any): 0.933 ROC-AUC, 87% landslides flagged")
set_bullet(slide, "TextBox 100", "Recognition / Awards (if any): None yet")

# --------------------------------------------------------- 06 FUTURE PLAN -
bullets_06 = [
    "Arduino ground-sensor node feeding the live dashboard.",
    "SMS gateway integration (logic ready, account pending).",
    "Validation partnership with GSI / state disaster mgmt.",
]
for shp_name, text in zip(["TextBox 123", "TextBox 125", "TextBox 127"], bullets_06):
    set_bullet(slide, shp_name, text)

# ------------------------------------------------------- 07 HOW IT WORKS --
# Input / Process / Solution are template-generic and fit their tiny
# (0.4in) boxes fine; "Outcome" alone wraps to "Outcom"/"e" at that width.
set_text(by_name(slide, "TextBox 151"), "Alert")

# ---------------------------------------------------------- 08 LOOKING FOR-
check(by_name(slide, "TextBox 157"), "Mentorship")
check(by_name(slide, "TextBox 158"), "Technology Partnership")
check(by_name(slide, "TextBox 158"), "Pilot / Deployment")

# --------------------------------------------------------------- QR code --
fit_picture(slide, "/Users/utkarsh/Desktop/ml/exhibition/qr_github.png", (3.28, 9.9, 1.4, 1.1))

# ------------------------------------------------------------- tagline ----
set_text(by_name(slide, "TextBox 195"), "#SPIREXA")

pres.save(OUT)
print("saved", OUT)
